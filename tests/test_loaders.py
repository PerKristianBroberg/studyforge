from pathlib import Path

import docx
import pptx

from studyforge.loaders import load
from studyforge.paths import classify


def meta_for(root: Path, rel: str):
    return classify(root, root / rel)


def write(root: Path, rel: str, text: str) -> Path:
    p = root / rel
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(text, encoding="utf-8")
    return p


def test_markdown_flattens_wikilinks(tmp_path):
    f = write(tmp_path, "UiO/INF1010/a.md", "See [[Quicksort|quick sort]] and [[Heaps]].")
    [doc] = load(f, meta_for(tmp_path, "UiO/INF1010/a.md"))
    assert doc.text == "See quick sort and Heaps."
    assert doc.id_ == "UiO/INF1010/a.md"
    assert doc.metadata == {"path": "UiO/INF1010/a.md", "uni": "UiO", "course": "INF1010", "doc_type": "notes"}


def test_code_keeps_text_and_language(tmp_path):
    src = "def f(x):\n    return [[x]]\n"
    f = write(tmp_path, "UiO/INF1010/code/f.py", src)
    [doc] = load(f, meta_for(tmp_path, "UiO/INF1010/code/f.py"))
    assert doc.text == src  # wikilink regex must not touch code
    assert doc.metadata["language"] == "python"
    assert doc.metadata["doc_type"] == "code"


def test_docx_is_single_document(tmp_path):
    f = tmp_path / "UiO/INF1010/notes.docx"
    f.parent.mkdir(parents=True)
    d = docx.Document()
    d.add_heading("Binary trees", level=1)
    d.add_paragraph("A binary tree node has at most two children.")
    d.save(f)
    [doc] = load(f, meta_for(tmp_path, "UiO/INF1010/notes.docx"))
    assert "at most two children" in doc.text
    assert "page" not in doc.metadata


def test_pptx_one_document_per_slide(tmp_path):
    f = tmp_path / "UiO/INF1010/slides.pptx"
    f.parent.mkdir(parents=True)
    deck = pptx.Presentation()
    for title in ["Stacks are LIFO", "Queues are FIFO"]:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
    deck.save(f)
    docs = load(f, meta_for(tmp_path, "UiO/INF1010/slides.pptx"))
    assert [d.metadata["page"] for d in docs] == [1, 2]
    assert [d.id_ for d in docs] == ["UiO/INF1010/slides.pptx#p1", "UiO/INF1010/slides.pptx#p2"]
    assert "LIFO" in docs[0].text and "FIFO" in docs[1].text


class _FakeDoc:
    def __init__(self, text):
        self.text, self.pages = text, {1: None}

    def export_to_markdown(self, page_no=None, image_placeholder=""):
        return self.text


class _FakeConverter:
    def __init__(self, text):
        self.text = text

    def convert(self, file):
        return type("Result", (), {"document": _FakeDoc(self.text)})()


def _pdf(tmp_path):
    f = tmp_path / "UiB/INF100/scan.pdf"
    f.parent.mkdir(parents=True)
    f.write_bytes(b"%PDF-1.4")
    return f


def test_scanned_pdf_falls_back_to_ocr(tmp_path, monkeypatch):
    from studyforge import loaders
    calls = []
    monkeypatch.setattr(loaders, "_converter",
                        lambda ocr=False: calls.append(ocr) or _FakeConverter("Haugen er et binært tre" if ocr else "  "))
    f = _pdf(tmp_path)
    [doc] = load(f, meta_for(tmp_path, "UiB/INF100/scan.pdf"))
    assert calls == [False, True]
    assert doc.text == "Haugen er et binært tre" and doc.metadata["page"] == 1


def test_text_pdf_does_not_run_ocr(tmp_path, monkeypatch):
    from studyforge import loaders
    calls = []
    monkeypatch.setattr(loaders, "_converter", lambda ocr=False: calls.append(ocr) or _FakeConverter("Real text"))
    f = _pdf(tmp_path)
    load(f, meta_for(tmp_path, "UiB/INF100/scan.pdf"))
    assert calls == [False]  # OCR is slow: only for PDFs without a text layer
